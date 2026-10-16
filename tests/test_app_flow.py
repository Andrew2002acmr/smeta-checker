"""Exercise the Streamlit audit and exports with a synthetic project and mocked API."""

import csv
import io
import json
import os
from pathlib import Path
import sys
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import fitz
import pandas as pd
import docx
from streamlit.testing.v1 import AppTest


class FakeClient:
    def __init__(self, **kwargs):
        self.responses = SimpleNamespace(create=self.create)

    def create(self, *, model, instructions, input, max_output_tokens, **_):
        prompt = input
        if "Проверяемая работа: Pump station" in prompt:
            raise RuntimeError("Simulated provider failure")
        data = {
            "work_found": True,
            "volume_confirmed": True,
            "unit_confirmed": True,
            "comment_confirmed": True,
            "found_volume": "45 m",
            "found_unit": "m",
            "found_page": "1",
            "evidence": [
                {"source_id": "p1q1", "kind": "work"},
                {"source_id": "p2q1", "kind": "volume"},
                {"source_id": "p2q1", "kind": "unit"},
            ],
            "discrepancy_details": "Confirmed with multi-page evidence",
        }
        return SimpleNamespace(output_text=json.dumps(data), usage=SimpleNamespace(total_tokens=100))

    def close(self):
        pass


class AppFlowTests(unittest.TestCase):
    def _run_audit(self, folder, header, rows, preparation_fault_on_call=None):
        with fitz.open() as pdf:
            for content in ("Copper cable specification. " * 20, "Copper cable 45 m. " * 20):
                page = pdf.new_page()
                page.insert_textbox(fitz.Rect(50, 50, 500, 600), content)
            pdf.save(folder / "project.pdf")
        with (folder / "vor.csv").open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.writer(stream, delimiter=";")
            writer.writerow(header)
            writer.writerows(rows)
        from pdf_text import PDFTextSource
        real_for_position = PDFTextSource.for_position
        state = {"calls": 0}

        def _maybe_flaky(source, *args, **kwargs):
            state["calls"] += 1
            if preparation_fault_on_call == state["calls"]:
                raise RuntimeError("Simulated preparation failure")
            return real_for_position(source, *args, **kwargs)

        with patch.dict(os.environ, {"YANDEX_CLOUD_API_KEY": "test-only", "YANDEX_CLOUD_FOLDER": "test-folder"}), \
                patch("openai.OpenAI", FakeClient), \
                patch("audit_tracing._configured_client", return_value=None), \
                patch("pdf_text.PDFTextSource.for_position", _maybe_flaky):
            at = AppTest.from_file(str(Path(__file__).resolve().parents[1] / "app.py"), default_timeout=30)
            at.run()
            self.assertEqual(len(at.exception), 0, str(at.exception))
            self.assertEqual([i for i in at.selectbox if i.label in ("Провайдер", "Модель")], [])
            folder_input = next(item for item in at.text_input if item.label == "Путь к папке проекта на диске")
            folder_input.set_value(str(folder)).run()
            self.assertEqual(len(at.exception), 0, str(at.exception))
            next(item for item in at.button if item.label == "🚀 Запустить аудит").click().run()
            self.assertEqual(len(at.exception), 0, str(at.exception))
            return at

    def test_audit_continues_after_failed_row_and_exports_actual_page(self):
        with tempfile.TemporaryDirectory(prefix="smeta-audit-test-") as folder:
            folder = Path(folder)
            at = self._run_audit(
                folder,
                ["№ п.п.", "Наименование работ, ресурсов, затрат по проекту",
                 "Объём работ / Количество", "Ед. изм.", "Ссылка на ПД",
                 "Наименование файла", "Номер страниц (через пробел)"],
                [[1, "Copper cable", 45, "m", "", "project.pdf", 1],
                 [2, "Pump station", 1, "шт", "", "project.pdf", 2],
                 [3, "Copper cable", 45, "m", "", "project.pdf", 2]])
            saved = at.session_state["audit_results"]["vor.csv"]
            table = saved["df"]
            self.assertEqual(table["Статус аудита"].tolist(), ["ОК", "Технический сбой", "Ошибка"])
            self.assertEqual(table["[ПД] Фактическая страница"].tolist(), ["Стр. 1, 2", "-", "Стр. 1, 2"])
            self.assertEqual(table["[ПД] Найденная единица"].tolist(), ["m", "-", "m"])
            self.assertEqual(table["[ПД] Найденные страницы"].tolist(), ["1, 2", "-", "1, 2"])
            self.assertTrue(all("Работа: да" in v for v in table["[ПД] Подтверждения (работа/объём/ед./прим.)"].tolist()[:1] + table["[ПД] Подтверждения (работа/объём/ед./прим.)"].tolist()[2:]))
            self.assertIn("Copper cable 45 m", table["[ПД] Доказательства (цитаты со страницами)"].iloc[0])
            self.assertIn("Copper cable 45 m", table["[ПД] Доказательства (цитаты со страницами)"].iloc[2])
            self.assertEqual(table["[ПД] Поиск завершён"].tolist(), ["да", "нет", "да"])
            tracking = saved["tracking"]
            self.assertEqual(tracking[0]["found_pages"], [1, 2])
            self.assertEqual(tracking[2]["found_pages"], [1, 2])
            self.assertEqual(tracking[0]["found_page"], 1)
            self.assertEqual(len(tracking[0]["evidence"]), 3)
            self.assertEqual(len(tracking[2]["evidence"]), 3)
            excel = pd.read_excel(io.BytesIO(saved["excel_bytes"]))
            self.assertEqual(excel["№"].tolist(), [3, 2, 1])
            self.assertEqual(excel["Статус аудита"].tolist(),
                             ["Обнаружено расхождение", "Проверка не выполнена", "ОК"])
            self.assertEqual(excel["Где проверить"].tolist(),
                             ["стр. 1, 2", "Место не установлено; ссылка ВОР: стр. 2", "стр. 1, 2"])
            self.assertFalse(any(str(col).startswith("[ПД]") for col in excel.columns))
            word = docx.Document(io.BytesIO(saved["word_bytes"]))
            self.assertEqual([r.cells[0].text for r in word.tables[0].rows[1:]], ["3", "2", "1"])
            self.assertIn("стр. 1, 2", word.tables[0].rows[1].cells[3].text)
            self.assertNotIn("Simulated provider failure", str(excel.values))
            self.assertNotIn("Simulated provider failure", word.tables[0].rows[2].cells[3].text)
            for pdf in at.session_state["pdf_pool"].values():
                pdf.close()

    def test_missing_explicit_pdf_never_binds_sole_file(self):
        with tempfile.TemporaryDirectory(prefix="smeta-audit-test-") as folder:
            folder = Path(folder)
            at = self._run_audit(
                folder,
                ["№ п.п.", "Наименование работ, ресурсов, затрат по проекту",
                 "Объём работ / Количество", "Ед. изм.", "Ссылка на ПД",
                 "Наименование файла", "Номер страниц (через пробел)"],
                [[1, "Copper cable", 45, "m", "", "ghost.pdf", 1]])
            saved = at.session_state["audit_results"]["vor.csv"]
            table = saved["df"]
            self.assertEqual(table["Статус аудита"].tolist(), ["Ошибка"])
            self.assertIn("ghost.pdf", table["Детали расхождений"].iloc[0])
            for pdf in at.session_state["pdf_pool"].values():
                pdf.close()

    def test_continuation_row_folds_refs_into_owner(self):
        with tempfile.TemporaryDirectory(prefix="smeta-audit-test-") as folder:
            folder = Path(folder)
            at = self._run_audit(
                folder,
                ["№ п.п.", "Наименование работ, ресурсов, затрат по проекту",
                 "Объём работ / Количество", "Ед. изм.", "Ссылка на ПД",
                 "Наименование файла", "Номер страниц (через пробел)"],
                [[1, "Copper cable", 45, "m", "", "project.pdf", 1],
                 ["", "", "", "", "", "project.pdf", 2]])
            saved = at.session_state["audit_results"]["vor.csv"]
            table = saved["df"]
            self.assertEqual(len(table), 1)
            self.assertEqual(table["Статус аудита"].tolist(), ["ОК"])
            for pdf in at.session_state["pdf_pool"].values():
                pdf.close()

    def test_distinct_documents_are_not_shown_ready_then_silently_partially_audited(self):
        with tempfile.TemporaryDirectory(prefix="smeta-audit-test-") as folder:
            folder = Path(folder)
            with fitz.open() as pdf:
                pdf.new_page()
                pdf.save(folder / "second.pdf")
            at = self._run_audit(
                folder,
                ["№ п.п.", "Наименование работ, ресурсов, затрат по проекту",
                 "Объём работ / Количество", "Ед. изм.", "Ссылка на ПД",
                 "Наименование файла", "Номер страниц (через пробел)"],
                [[1, "Copper cable", 45, "m", "", "project.pdf;second.pdf", 1]])
            try:
                table = at.session_state["audit_results"]["vor.csv"]["df"]
                self.assertEqual(table["Статус аудита"].tolist(), ["Ошибка"])
                self.assertIn("несколько PDF", table["Детали расхождений"].iloc[0])
                preflight = next(frame.value for frame in at.dataframe
                                 if "Статус привязки" in frame.value.columns)
                self.assertTrue(preflight["Статус привязки"].iloc[0].startswith("🔴"))
            finally:
                for pdf in at.session_state["pdf_pool"].values():
                    pdf.close()

    def test_partly_invalid_page_list_cannot_be_ok(self):
        with tempfile.TemporaryDirectory(prefix="smeta-audit-test-") as folder:
            at = self._run_audit(
                Path(folder),
                ["№ п.п.", "Наименование работ, ресурсов, затрат по проекту",
                 "Объём работ / Количество", "Ед. изм.", "Ссылка на ПД",
                 "Наименование файла", "Номер страниц (через пробел)"],
                [[1, "Copper cable", 45, "m", "", "project.pdf", "1 999"]])
            try:
                table = at.session_state["audit_results"]["vor.csv"]["df"]
                self.assertEqual(table["Статус аудита"].tolist(), ["Ошибка"])
                self.assertIn("999", table["Детали расхождений"].iloc[0])
                self.assertEqual(table["[ПД] Найденные страницы"].tolist(), ["1, 2"])
            finally:
                for pdf in at.session_state["pdf_pool"].values():
                    pdf.close()

    def test_audit_reads_vor_page_link_column(self):
        with tempfile.TemporaryDirectory(prefix="smeta-audit-test-") as folder:
            folder = Path(folder)
            at = self._run_audit(
                folder,
                ["№ п.п.", "Наименование работ, ресурсов, затрат по проекту",
                 "Объём работ / Количество", "Ед. изм.", "Ссылка на чертежи",
                 "Наименование файла", "[ВОР] Ссылка на страницу"],
                [[1, "Copper cable", 45, "m", "", "project.pdf", 1]])
            saved = at.session_state["audit_results"]["vor.csv"]
            table = saved["df"]
            self.assertEqual(table["Статус аудита"].tolist(), ["ОК"])
            self.assertNotIn("Страница не указана", table["Детали расхождений"].iloc[0])
            self.assertEqual(table["[ВОР] Ссылка на страницу"].tolist(), [1])
            for pdf in at.session_state["pdf_pool"].values():
                pdf.close()

    def test_audit_rerun_overwrites_stale_results(self):
        with tempfile.TemporaryDirectory(prefix="smeta-audit-test-") as folder:
            folder = Path(folder)
            at = self._run_audit(
                folder,
                ["№ п.п.", "Наименование работ, ресурсов, затрат по проекту",
                 "Объём работ / Количество", "Ед. изм.", "Ссылка на чертежи",
                 "Наименование файла", "[ВОР] Ссылка на страницу",
                 "Статус аудита", "Детали расхождений",
                 "[ПД] Найденный объем (нейросеть)", "[ПД] Фактическая страница"],
                [[1, "Copper cable", 45, "m", "", "project.pdf", 1,
                  "Технический сбой", "stale", "-", "-"]])
            saved = at.session_state["audit_results"]["vor.csv"]
            table = saved["df"]
            self.assertEqual(list(table.columns).count("Статус аудита"), 1)
            self.assertEqual(table["Статус аудита"].tolist(), ["ОК"])
            for pdf in at.session_state["pdf_pool"].values():
                pdf.close()


    def test_broken_preparation_does_not_abort_other_positions(self):
        with tempfile.TemporaryDirectory(prefix="smeta-audit-test-") as folder:
            at = self._run_audit(
                Path(folder),
                ["№ п.п.", "Наименование работ, ресурсов, затрат по проекту",
                 "Объём работ / Количество", "Ед. изм.", "Ссылка на ПД",
                 "Наименование файла", "Номер страниц (через пробел)"],
                [[1, "Copper cable", 45, "m", "", "project.pdf", 1],
                 [2, "Pump station", 1, "шт", "", "project.pdf", 2],
                 [3, "Copper cable", 45, "m", "", "project.pdf", 2]],
                preparation_fault_on_call=2)
            try:
                saved = at.session_state["audit_results"]["vor.csv"]
                table = saved["df"]
                self.assertEqual(table["Статус аудита"].tolist(),
                                 ["ОК", "Технический сбой", "Ошибка"])
                self.assertIn("Simulated preparation failure",
                              table["Детали расхождений"].iloc[1])
                self.assertEqual(table["[ПД] Запросов"].iloc[1], 0)
                self.assertEqual(table["[ПД] Поиск завершён"].iloc[1], "нет")
                self.assertEqual(saved["tracking"][1]["work_name"], "Pump station")
                self.assertIn("trace_id", saved["tracking"][1])
                excel = pd.read_excel(io.BytesIO(saved["excel_bytes"]))
                statuses = dict(zip(excel["№"], excel["Статус аудита"]))
                self.assertEqual(statuses, {1: "ОК", 2: "Проверка не выполнена",
                                            3: "Обнаружено расхождение"})
            finally:
                for pdf in at.session_state["pdf_pool"].values():
                    pdf.close()


if __name__ == "__main__":
    unittest.main()
