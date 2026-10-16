"""User-visible ordering, evidence boundaries and document navigation."""

import copy
import io
from pathlib import Path
import sys
import unittest
from urllib.parse import unquote

import docx
import openpyxl
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from audit_reports import (DISCREPANCY, MANUAL, NOT_CHECKED, build_user_report,
                           generate_excel_report, generate_word_report)


class ReportTests(unittest.TestCase):
    def make_report(self, cases):
        frame = pd.DataFrame([{"№ п.п.": i + 1, "Статус аудита": status,
                               "Детали расхождений": "found_volume=null, search_completed=true; OCR; response_status=failed",
                               "Идентификатор": "internal-id", "Тип позиции": "Работа",
                               "Формула расчета объемов работ и расхода материалов, потребности ресурсов": "10*12",
                               "Дополнительная информация (комментарий).": "Скрыто",
                               "[ВОР] Заявленный объем": "120", "[ВОР] Ед. изм.": "м³",
                               "[ПД] Запросов": 6} for i, (status, _) in enumerate(cases)], index=range(17, 17 + len(cases)))
        sources, tasks = [], []
        for status, overrides in cases:
            source = {"pdf_name": "кр.pdf", "found_page": 18, "found_pages": [18, 19],
                      "vor_pages": [18], "work_found": True, "volume_confirmed": True,
                      "unit_confirmed": True, "comment_confirmed": True,
                      "found_volume": "120", "found_unit": "м³", "evidence": [], "requests": 6}
            source.update(overrides)
            sources.append(source)
            tasks.append({"w_name": "Бетонные работы — полное наименование без обрезки на 45 символах",
                          "w_vol": "120", "w_unit": "м³", "w_comment": "", "w_page": "18",
                          "target_indices": [17], "loc_desc": "Стр. 18", "file_refs": r"C:\project\КР.pdf"})
        before_frame, before_sources = frame.copy(deep=True), copy.deepcopy(sources)
        report = build_user_report(frame, sources, tasks)
        pd.testing.assert_frame_equal(frame, before_frame)
        self.assertEqual(sources, before_sources)
        return report

    def test_priorities_stable_order_and_missing_proof_not_a_discrepancy(self):
        report = self.make_report([
            ("ОК", {}),
            ("Неполная проверка", {"found_volume": "-", "volume_confirmed": False}),
            ("Ошибка", {"found_volume": "95", "volume_confirmed": False}),
            ("Технический сбой", {"work_found": False, "found_pages": [], "found_page": None}),
            ("Ошибка", {"work_found": False, "found_pages": [], "found_page": None}),
            ("ОК", {})])
        self.assertEqual(report["№"].tolist(), [3, 2, 5, 4, 1, 6])
        self.assertEqual(report["Статус аудита"].tolist(), [DISCREPANCY, MANUAL, MANUAL, NOT_CHECKED, "ОК", "ОК"])
        mismatch = report.iloc[0]
        self.assertIn("120 м³", mismatch["Что не так"])
        self.assertIn("95 м³", mismatch["Что не так"])
        self.assertIn("25 м³", mismatch["Что не так"])
        self.assertEqual(mismatch["Наименование файла"], "КР.pdf")
        self.assertEqual(mismatch["Где проверить"], "стр. 18, 19")
        for text in report["Что не так"]:
            self.assertNotIn("found_volume", text)
            self.assertNotIn("OCR", text)
        self.assertNotIn("Идентификатор", report.columns)
        self.assertIn("не доказывает", report.loc[report["№"] == 5, "Что не так"].iloc[0])

    def test_reference_and_unit_discrepancies_not_hidden_by_incomplete_audit(self):
        report = self.make_report([
            ("Ошибка", {"found_page": 19}),
            ("Неполная проверка", {"unit_confirmed": False, "found_unit": "м²"}),
            ("Неполная проверка", {"comment_confirmed": False}),
            ("Ошибка", {"volume_confirmed": False, "found_volume": "120 м³"})])
        self.assertEqual(report["Статус аудита"].tolist(), [DISCREPANCY, DISCREPANCY, MANUAL, MANUAL])
        self.assertIn("вне страниц", report.iloc[0]["Что не так"])
        self.assertIn("не полностью", report.iloc[1]["Что не так"])
        self.assertNotIn("Разница", report.iloc[3]["Что не так"])

    def test_derived_volume_is_not_presented_as_printed_quantity(self):
        report = self.make_report([("Ошибка", {
            "volume_confirmed": False, "found_volume": "95",
            "volume_calculation": {"area": {"value": "950", "unit": "м2"},
                                   "thickness": {"value": "0,1", "unit": "м"}}})])
        description = report.iloc[0]["Что не так"]
        self.assertIn("по расчёту из ПД", description)
        self.assertIn("950 м2 × 0,1 м", description)
        self.assertIn("25 м³", description)

    def test_excel_colors_literal_names_and_word_navigation(self):
        report = self.make_report([
            ("ОК", {}),
            ("Ошибка", {"found_volume": "95", "volume_confirmed": False}),
            ("Неполная проверка", {"volume_confirmed": False, "found_volume": "-"}),
            ("Технический сбой", {"work_found": False, "found_pages": [], "found_page": None})])
        report.loc[0, "Наименование работ"] = "=1+1"
        excel = openpyxl.load_workbook(io.BytesIO(generate_excel_report(report)))
        sheet = excel.active
        self.assertEqual([sheet.cell(i, 3).fill.fgColor.rgb for i in range(2, 6)],
                         ["00FEE2E2", "00FEF3C7", "00E5E7EB", "00DCFCE7"])
        self.assertEqual(sheet.cell(2, 2).data_type, "s")
        self.assertEqual(unquote(sheet.cell(2, 5).hyperlink.target), "КР.pdf#page=18")
        word = docx.Document(io.BytesIO(generate_word_report(report, {})))
        table = word.tables[0]
        self.assertEqual(len(table.columns), 4)
        self.assertEqual([r.cells[0].text for r in table.rows[1:]], ["2", "3", "4", "1"])
        self.assertIn("120 м³", table.rows[1].cells[3].text)
        self.assertIn("стр. 18, 19", table.rows[1].cells[3].text)
        self.assertEqual(table.rows[1].cells[1].text, "=1+1")
        targets = [unquote(rel.target_ref) for rel in word.part.rels.values() if rel.reltype.endswith("/hyperlink")]
        self.assertIn("КР.pdf#page=18", targets)
        self.assertNotIn("found_volume", " ".join(c.text for r in table.rows for c in r.cells))

    def test_missing_or_multiple_documents_have_no_invented_source_link(self):
        frame = pd.DataFrame([{"№": 1, "Статус аудита": "Ошибка"}])
        task = {"w_name": "Работа", "file_refs": "/storage/КР.pdf;..\\АР.pdf", "bind_error": "internal error"}
        report = build_user_report(frame, [{"pdf_name": "", "found_pages": []}], [task])
        self.assertEqual(report.iloc[0]["Наименование файла"], "КР.pdf\nАР.pdf")
        self.assertEqual(report.iloc[0]["Статус аудита"], NOT_CHECKED)
        sheet = openpyxl.load_workbook(io.BytesIO(generate_excel_report(report))).active
        self.assertIsNone(sheet.cell(2, 5).hyperlink)
        self.assertEqual(report.iloc[0]["Где проверить"], "Место не установлено")


if __name__ == "__main__":
    unittest.main()
