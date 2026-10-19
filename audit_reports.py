"""User-facing audit exports; raw audit data and row order are never mutated."""

import io
import re
from decimal import Decimal
from urllib.parse import quote

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.shared import Inches, Pt
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from audit_core import INCOMPLETE, TECHNICAL_ERROR, format_ru
from audit_models import AuditResult, VorPosition
from pdf_binding import split_pdf_refs

DISCREPANCY = "Обнаружено расхождение"
MANUAL = "Требуется ручная проверка"
NOT_CHECKED = "Проверка не выполнена"
STATUS_ORDER = {DISCREPANCY: 0, MANUAL: 1, NOT_CHECKED: 2, "ОК": 3}
STATUS_COLORS = {DISCREPANCY: "FEE2E2", MANUAL: "FEF3C7",
                 NOT_CHECKED: "E5E7EB", "ОК": "DCFCE7"}
REPORT_COLUMNS = ["№", "Наименование работ", "Статус аудита", "Что не так",
                  "Наименование файла", "Где проверить"]


def _text(value):
    text = str(value).strip() if value is not None else ""
    return "" if text.lower() in {"", "nan", "none", "null", "-"} else text


def _basename(value):
    return value.replace("\\", "/").rsplit("/", 1)[-1]


def _quantity(value, unit):
    value, unit = format_ru(_text(value)), _text(unit)
    return value if not unit or value.endswith(unit) else f"{value} {unit}".strip()


def _number(value, unit):
    value = _text(value)
    if unit and value.endswith(unit):
        value = value[:-len(unit)].strip()
    if re.fullmatch(r"[+-]?\d+(?:[.,]\d+)?", value):
        return Decimal(value.replace(",", "."))
    return None


def _as_task_dict(task):
    if isinstance(task, VorPosition):
        return task.legacy_task()
    return task


def _as_source_dict(source):
    if isinstance(source, AuditResult):
        data = source.to_legacy()
        data["found_page"] = (
            int(source.found_page) if source.found_page.isdigit() else None
        )
        return data
    return source


def build_user_report(df, tracking, tasks):
    """Project explicit presentation fields; align tracking by position, not df index."""
    tracking = [_as_source_dict(source) for source in tracking]
    tasks = [_as_task_dict(task) for task in tasks]
    if len(df) != len(tracking) or len(df) != len(tasks):
        raise ValueError("Данные позиций и источников отчёта не согласованы")
    records = []
    for (_, row), source, task in zip(df.iterrows(), tracking, tasks):
        files = [_basename(ref) for ref in split_pdf_refs(task.get("file_refs", ""))]
        if source.get("pdf_name"):
            resolved = _basename(source["pdf_name"])
            files = [next((name for name in files if name.casefold() == resolved.casefold()), resolved)]
        files = list(dict.fromkeys(files))
        pages = sorted(set(source.get("found_pages") or []))
        location = ("стр. " + ", ".join(map(str, pages))) if pages else "Место не установлено"
        if not pages and source.get("vor_pages"):
            location += "; ссылка ВОР: стр. " + ", ".join(map(str, source["vor_pages"]))
        consumption = source.get("consumption") or {}
        if consumption.get("schedule_page") and consumption["schedule_page"] not in pages:
            extra = f"стр. {consumption['schedule_page']}"
            location = extra if location == "Место не установлено" else location + "; " + extra
        remarks, notes = [], []
        raw_status = row.get("Статус аудита")
        status = MANUAL
        expected_volume, expected_unit = _text(task.get("w_vol")), _text(task.get("w_unit"))
        actual_volume, actual_unit = _text(source.get("found_volume")), _text(source.get("found_unit"))
        if task.get("bind_error"):
            status = NOT_CHECKED
            remarks.append("Непонятно, к какому файлу относится позиция: имя в ВОР не совпало ни с одним PDF.")
        elif not source.get("work_found"):
            if raw_status == TECHNICAL_ERROR:
                status = NOT_CHECKED
                remarks.append("Проверка сорвалась: документ не дочитан или модель не ответила.")
            else:
                remarks.append("Работа в проекте не нашлась — но это ещё не значит, что её там нет.")
        else:
            before = _number(expected_volume, expected_unit)
            after = _number(actual_volume, actual_unit)
            same_unit = expected_unit.casefold().replace("³", "3").replace("²", "2") == actual_unit.casefold().replace("³", "3").replace("²", "2")
            calculation = source.get("volume_calculation")
            if consumption.get("matches") is False:
                status = DISCREPANCY
                chain = consumption.get("stock_chain") or {}
                if chain:
                    remarks.append(
                        f"Потребность в материале с запасом не сошлась: в ВОР {chain['material_volume']} {expected_unit} "
                        f"(работа {chain['work_volume']} × {chain['factor']}, позиция №{chain['material_number']}), "
                        f"а в спецификации — {_quantity(consumption.get('schedule_volume'), expected_unit)} "
                        f"(стр. {consumption.get('schedule_page')}, {consumption.get('schedule_item')}). "
                        f"Разница — {_quantity(consumption.get('delta'), expected_unit)} "
                        f"({consumption.get('delta_percent')}%).")
                else:
                    remarks.append(
                        f"В ВОР {_quantity(expected_volume, expected_unit)}, "
                        f"а в спецификации — {_quantity(consumption.get('schedule_volume'), expected_unit)} "
                        f"(стр. {consumption.get('schedule_page')}, {consumption.get('schedule_item')}). "
                        f"Разница — {_quantity(consumption.get('delta'), expected_unit)} "
                        f"({consumption.get('delta_percent')}%). "
                        f"Норма в ВОР — {consumption.get('rate')} {consumption.get('unit')}, "
                        f"а из спецификации получается {consumption.get('derived_rate')} {consumption.get('unit')}.")
            elif consumption.get("matches") is True:
                chain = consumption.get("stock_chain") or {}
                if chain:
                    notes.append(
                        f"Потребность в материале сходится: {chain['work_volume']} × {chain['factor']} = "
                        f"{chain['material_volume']} {expected_unit} (позиция №{chain['material_number']}), "
                        f"в спецификации — {consumption.get('schedule_volume')} {expected_unit} "
                        f"(стр. {consumption.get('schedule_page')}).")
                else:
                    notes.append(
                        f"Норма {consumption.get('rate')} {consumption.get('unit')} сошлась "
                        f"со спецификацией (стр. {consumption.get('schedule_page')}).")
            elif not source.get("volume_confirmed"):
                if before is not None and after is not None and before != after and same_unit:
                    status = DISCREPANCY
                    delta = format_ru(abs(before - after))
                    origin = "по расчёту" if calculation else "в проекте"
                    remarks.append(
                        f"В ВОР {_quantity(expected_volume, expected_unit)}, {origin} — "
                        f"{_quantity(actual_volume, actual_unit)}. Разница — {delta} {expected_unit}.")
                elif source.get("quantity_basis") == "consumption":
                    remarks.append(
                        f"Объём {_quantity(expected_volume, expected_unit)} посчитан по норме расхода — "
                        "в проекте готового числа нет, сверяйтесь со спецификацией.")
                elif (before is not None and after is not None and before == after and same_unit):
                    notes.append(
                        f"Число {_quantity(expected_volume, expected_unit)} в проекте есть, "
                        "но как объём позиции не подтверждено — нужна сверка.")
                else:
                    remarks.append(f"Объём {_quantity(expected_volume, expected_unit)} в проекте не нашёлся" + (f" — есть {_quantity(actual_volume, actual_unit)}." if actual_volume else "."))
            if not source.get("unit_confirmed"):
                if actual_unit and expected_unit and not same_unit and "/" not in actual_unit and "/" not in expected_unit:
                    status = DISCREPANCY
                    remarks.append(f"В ВОР единица «{expected_unit}», а в проекте — «{actual_unit}».")
                else:
                    notes.append(f"Единица «{expected_unit}» в проекте прямо не подтверждена.")
            technology = source.get("technology") or {}
            tech_applies = str(source.get("position_type", "Работа") or "Работа").casefold() == "работа"
            if technology.get("status") == "different" and tech_applies:
                quantities_ok = bool(source.get("volume_confirmed") and source.get("unit_confirmed"))
                if quantities_ok:
                    notes.append(f"Способ выполнения различается: в ВОР — «{technology.get('expected', '')}», "
                                 f"а в проекте — «{technology.get('actual', '')}»; нужна инженерная сверка.")
                else:
                    status = DISCREPANCY
                    remarks.append(f"В ВОР — «{technology.get('expected', '')}», "
                                   f"а в проекте — «{technology.get('actual', '')}».")
            elif technology.get("status") == "unconfirmed" and tech_applies:
                cross = [p for p in (technology.get("cross_pages") or [])
                         if isinstance(p, int)]
                if cross:
                    pages = ", ".join(map(str, cross))
                    notes.append(f"Способ выполнения в проекте не нашёлся — сама работа есть. "
                                 f"Формулировка ВОР также встречается на стр. {pages} — сверьте там.")
                else:
                    notes.append("Способ выполнения в проекте не нашёлся — сама работа есть.")
            workbook = source.get("workbook_calculation") or {}
            if workbook.get("status") in {"error", "unsupported"}:
                remarks.append("Формула в Excel не сошлась: " + str(workbook.get("error", workbook["status"])) + ".")
            for note in source.get("workbook_notes") or []:
                kind = (note or {}).get("kind", "")
                message = str((note or {}).get("message", "")).strip()
                if not message:
                    continue
                if kind in {"scheme-divisor-mismatch", "length-operand-conflicts"}:
                    notes.append("Вопрос к исходным данным: " + message + ".")
                else:
                    status = DISCREPANCY
                    remarks.append("В самой ведомости противоречие: " + message + ".")
            for code in consumption.get("uncertainties") or []:
                notes.append({"schedule-row-unlinked": "Строку спецификации под эту норму выбрать однозначно не вышло.",
                             "scope-unresolved": "Непонятно, к какому участку относится норма.",
                             "area-unconfirmed": "Площадь под нормой не подтверждена.",
                             "rate-differs-from-schedule": "Норма ВОР не совпадает с нормой из спецификации.",
                             "claimed-volume-ambiguous": "Объём в ВОР записан неоднозначно.",
                             "schedule-volume-missing": "В спецификации нет количества.",
                             "schedule-text-unavailable": "Текст спецификации недоступен.",
                             "rate-justification-open": "Обоснованность нормы расхода в ПД не подтверждена.",
                             "chain-arithmetic-failed": "Цепочка «работа × запас = материал» не сошлась.",
                             "chain-mismatch": "Цепочка «работа × запас = материал» не сошлась.",
                             }.get(code, code))
            if not source.get("comment_confirmed"):
                notes.append("Примечание из ВОР в проекте не нашлось.")
            found_page = source.get("found_page")
            if found_page and source.get("vor_pages") and found_page not in source["vor_pages"]:
                if any(abs(found_page - page) == 1 for page in source["vor_pages"]):
                    notes.append(f"Нашлось на стр. {found_page} — рядом со ссылкой ВОР.")
                else:
                    notes.append("Нашлось не на тех страницах, что указаны в ВОР "
                                 f"(стр. {', '.join(map(str, source['vor_pages']))}) — "
                                 "навигационная пометка, не противоречие количества.")
            if not task.get("target_indices") or "недопустимые" in task.get("loc_desc", ""):
                notes.append("В ВОР нет годной ссылки на страницы: " + (_text(task.get("w_page")) or "не указана") + " — "
                             "дефект ведомости, расхождение с ПД не доказано.")
            if raw_status in {INCOMPLETE, TECHNICAL_ERROR}:
                notes.append("Проверка неполная — часть сведений осталась за кадром.")
            if not remarks and not notes and raw_status == "ОК":
                status = "ОК"
            elif not remarks and not notes:
                notes.append("Нужно глянуть глазами: соответствие не очевидно.")
            elif not remarks and raw_status == "ОК":
                status = "ОК"
        records.append(dict(zip(REPORT_COLUMNS, [row.iloc[0], task["w_name"], status,
            " ".join(remarks + notes) if remarks or notes else "Расхождений не обнаружено.",
            "\n".join(files) or "Не указан", location])))
    report = pd.DataFrame(records, columns=REPORT_COLUMNS)
    return report.sort_values("Статус аудита", key=lambda col: col.map(STATUS_ORDER), kind="stable").reset_index(drop=True)


def _source_link(row):
    name = row["Наименование файла"]
    if name == "Не указан" or "\n" in name:
        return None
    match = re.match(r"стр\. (\d+)", row["Где проверить"])
    return quote(name) + (f"#page={match[1]}" if match else "")


def generate_excel_report(report):
    out = io.BytesIO()
    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        report.to_excel(writer, index=False, sheet_name="Аудит")
        sheet = writer.sheets["Аудит"]
        side = Side(style="thin", color="D3D3D3")
        border = Border(left=side, right=side, top=side, bottom=side)
        for cells in sheet:
            for cell in cells:
                # All document-supplied text is literal, never an Excel formula.
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                cell.border = border
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.font = Font(color="FFFFFF", bold=True, size=10)
        for i, (_, row) in enumerate(report.iterrows(), 2):
            sheet.cell(i, 3).fill = PatternFill("solid", fgColor=STATUS_COLORS[row["Статус аудита"]])
            link = _source_link(row)
            if link:
                sheet.cell(i, 5).hyperlink = link
                sheet.cell(i, 5).font = Font(color="0563C1", underline="single")
        for i, width in enumerate((8, 48, 28, 65, 32, 35), 1):
            sheet.column_dimensions[get_column_letter(i)].width = width
        sheet.freeze_panes = "C2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.print_title_rows = "1:1"
    return out.getvalue()


def generate_word_report(report, meta):
    doc = docx.Document()
    for section in doc.sections:
        section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(0.6)
    doc.styles["Normal"].font.size = Pt(9)
    doc.add_heading("Результаты аудита ВОР и ПД", 0)
    doc.add_paragraph(f"Объект: {meta.get('object_name', 'Не указано')}\nОснование: {meta.get('basis_code', 'ПД')}")
    counts = report["Статус аудита"].value_counts()
    doc.add_paragraph(f"Всего позиций: {len(report)}. " + "; ".join(f"{status}: {counts.get(status, 0)}" for status in STATUS_ORDER))
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ("№", "Наименование работ", "Статус аудита", "Результат и действия")):
        cell.text = text
        cell.paragraphs[0].runs[0].bold = True
    repeat = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for _, row in report.iterrows():
        cells = table.add_row().cells
        cells[0].text = str(row["№"])
        cells[1].text = str(row["Наименование работ"])
        cells[2].text = row["Статус аудита"]
        cells[2]._tc.get_or_add_tcPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{STATUS_COLORS[row["Статус аудита"]]}"/>'))
        cells[3].text = row["Что не так"]
        paragraph = cells[3].add_paragraph("Где проверить: ")
        link = _source_link(row)
        if link:
            hyperlink = OxmlElement("w:hyperlink")
            hyperlink.set(qn("r:id"), doc.part.relate_to(link, RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
            run = OxmlElement("w:r")
            properties = OxmlElement("w:rPr")
            style = OxmlElement("w:u")
            style.set(qn("w:val"), "single")
            properties.append(style)
            run.append(properties)
            text = OxmlElement("w:t")
            text.text = row["Наименование файла"]
            run.append(text)
            hyperlink.append(run)
            paragraph._p.append(hyperlink)
        else:
            paragraph.add_run(row["Наименование файла"])
        paragraph.add_run("; " + row["Где проверить"] + ".")
    for word_row in table.rows:
        for cell, width in zip(word_row.cells, (0.4, 1.5, 1.2, 3.9)):
            cell.width = Inches(width)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
